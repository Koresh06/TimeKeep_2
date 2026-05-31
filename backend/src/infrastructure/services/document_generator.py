from dataclasses import dataclass, field
from pathlib import Path
from docx import Document
from docx.shared import Pt
from io import BytesIO

from src.application.dtos.document import DocumentDTO
from src.domain.interfaces.services.document_generator import IDocumentGenerator

TEMPLATE_PATH = Path("src/infrastructure/services/templates/report_template.docx")


@dataclass
class DocumentGenerator(IDocumentGenerator):
    template_path: Path = field(default_factory=lambda: TEMPLATE_PATH)

    async def generate_day_off_document(
        self,
        dto: DocumentDTO,
    ) -> bytes:
        doc = Document(str(self.template_path))

        placeholders = {
            "organization_position_boss": dto.boss_position,
            "name_organization": dto.organization_name_genitive,
            "organization_rank_boss": dto.boss_rank,
            "organization_name_boss": dto.boss_full_name,
            "date_report": dto.date_report,
            "date_day_off": dto.date_day_off,
            "info_overtimes": dto.info_overtimes,
            "position_user": dto.position_user,
            "department_user": dto.department_user,
            "rank_user": dto.rank_user,
            "full_name_user": dto.full_name_user,
        }

        for paragraph in doc.paragraphs:
            for key, value in placeholders.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, str(value))
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(14)

        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.read()
