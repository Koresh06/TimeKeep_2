from dataclasses import dataclass, field
from pathlib import Path
from docx import Document
from docx.shared import Pt
from io import BytesIO

from src.domain.entities.day_off import DayOff
from src.domain.entities.organization import Organization
from src.domain.entities.user import User
from src.domain.interfaces.services.document_generator import IDocumentGenerator


TEMPLATE_PATH = Path("src/infrastructure/services/templates/report_template.docx")


@dataclass
class DocumentGenerator(IDocumentGenerator):
    template_path: Path = field(default_factory=lambda: TEMPLATE_PATH)

    async def generate_day_off_document(
        self,
        day_off: DayOff,
        user: User,
        organization: Organization,
        boss: User,
    ) -> bytes:
        doc = Document(str(self.template_path))

        placeholders = {
            "organization_position_boss": boss.position,
            "name_organization": organization.name,
            "organization_rank_boss": boss.rank,
            "organization_name_boss": f"{boss.surname} {boss.first_name[0]}.{boss.patronymic[0]}.",
            "date_report": day_off.created_at.strftime("%d.%m.%Y"),
            "date_day_off": day_off.date_.strftime("%d.%m.%Y"),
            "info_overtimes": day_off.format_overtimes_for_document(),
            "position_user": user.position,
            "rank_user": user.rank,
            "full_name_user": f"{user.surname} {user.first_name[0]}.{user.patronymic[0]}.",
        }

        for paragraph in doc.paragraphs:
            for key, value in placeholders.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, str(value))
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(14)

        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.read()